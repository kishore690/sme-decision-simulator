import os
import re
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
import numpy as np
import pandas as pd
from business import get_business_metrics
from ai import ask_ai, ask_business_question
from simulator import simulate_hiring
from price_simulator import simulate_price_change
from marketing_simulator import simulate_marketing
from loan_simulator import simulate_loan
from inventory_simulator import simulate_inventory
from expansion_simulator import simulate_expansion
from reports import (
    REPORT_TYPES,
    build_report_data,
    generate_all_reports_zip,
    generate_csv_report,
    generate_pdf_report,
)
app = FastAPI(title="SME Decision Simulator API", version="1.0.0")

# Store uploaded sales data in memory
uploaded_data = None
uploaded_business_context = None

# Allow React frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in os.getenv(
        "CORS_ORIGINS", "http://localhost:5173,http://127.0.0.1:5173"
    ).split(",") if origin.strip()],
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type"],
)

REQUIRED_COLUMNS = ["Date", "Product", "Quantity", "Revenue"]
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", str(50 * 1024 * 1024)))
MAX_OFFICE_EXPANDED_BYTES = 250 * 1024 * 1024
MAX_UPLOAD_ROWS = 1_000_000
SUPPORTED_UPLOAD_EXTENSIONS = {
    ".csv", ".tsv", ".txt", ".json", ".xlsx", ".xls", ".xlsm", ".docx", ".pdf"
}

SALES_COLUMN_ALIASES = {
    "Date": {"date", "orderdate", "invoicedate", "transactiondate", "saledate", "purchasedate", "datetime"},
    "Product": {"product", "productname", "item", "itemname", "description", "itemdescription", "stockdescription"},
    "Quantity": {"quantity", "qty", "units", "unitssold", "itemquantity"},
    "Revenue": {"revenue", "sales", "total", "amount", "linetotal", "salesamount", "totalamount", "netamount", "extendedprice"},
    "UnitPrice": {"unitprice", "priceperunit", "price", "rate"},
    "OrderID": {"orderid", "ordernumber", "invoiceno", "invoicenumber", "transactionid", "transactionnumber", "receiptid"},
}


def _column_token(value):
    return re.sub(r"[^a-z0-9]", "", str(value).casefold())


def _column_mapping(columns):
    source_by_token = {_column_token(column): column for column in columns}
    mapping = {}
    for canonical, aliases in SALES_COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in source_by_token:
                mapping[canonical] = source_by_token[alias]
                break
    return mapping


def _has_sales_headers(columns):
    mapping = _column_mapping(columns)
    return all(name in mapping for name in ("Date", "Product", "Quantity")) and (
        "Revenue" in mapping or "UnitPrice" in mapping
    )


def _validate_office_archive(raw):
    try:
        with ZipFile(BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 10_000 or sum(member.file_size for member in members) > MAX_OFFICE_EXPANDED_BYTES:
                raise HTTPException(status_code=413, detail="The expanded Office document is too large to process safely.")
    except BadZipFile as exc:
        raise HTTPException(status_code=422, detail="This Office file is damaged or is not a valid Excel/Word document.") from exc


def _frame_from_table(rows):
    cleaned = [[str(cell).strip() if cell is not None else "" for cell in row] for row in rows]
    cleaned = [row for row in cleaned if any(value for value in row)]
    if len(cleaned) < 2:
        return None
    headers = cleaned[0]
    if not _has_sales_headers(headers):
        return None
    width = len(headers)
    return pd.DataFrame([row[:width] + [None] * max(0, width - len(row)) for row in cleaned[1:]], columns=headers)


def _read_sales_file(extension, raw):
    if extension == ".csv":
        return pd.read_csv(BytesIO(raw), sep=None, engine="python", encoding="utf-8-sig")
    if extension == ".tsv":
        return pd.read_csv(BytesIO(raw), sep="\t", encoding="utf-8-sig")
    if extension == ".txt":
        return pd.read_csv(BytesIO(raw), sep=None, engine="python", encoding="utf-8-sig")
    if extension == ".json":
        return pd.read_json(BytesIO(raw))
    if extension in {".xlsx", ".xls", ".xlsm"}:
        if extension == ".xls":
            return pd.read_excel(BytesIO(raw), engine="xlrd")

        _validate_office_archive(raw)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                selected_columns = None
                records = None
                for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if selected_columns is None:
                        if row_number > 50:
                            break
                        headers = [str(value).strip() if value is not None else "" for value in values]
                        if not _has_sales_headers(headers):
                            continue
                        mapping = _column_mapping(headers)
                        names = ["Date", "Product", "Quantity"]
                        names.append("Revenue" if "Revenue" in mapping else "UnitPrice")
                        if "OrderID" in mapping:
                            names.append("OrderID")
                        selected_columns = [(name, headers.index(mapping[name])) for name in names]
                        records = {name: [] for name in names}
                        continue

                    if not any(value is not None for value in values):
                        continue
                    for name, index in selected_columns:
                        records[name].append(values[index] if index < len(values) else None)
                    if len(records["Date"]) > MAX_UPLOAD_ROWS:
                        raise HTTPException(status_code=422, detail=f"Sales files cannot contain more than {MAX_UPLOAD_ROWS:,} rows.")

                if records is not None:
                    return pd.DataFrame(records)
        finally:
            workbook.close()
        raise HTTPException(status_code=422, detail="The workbook contains no worksheet with recognizable sales columns.")
    if extension == ".docx":
        _validate_office_archive(raw)
        from docx import Document

        document = Document(BytesIO(raw))
        for table in document.tables:
            frame = _frame_from_table([[cell.text for cell in row.cells] for row in table.rows])
            if frame is not None:
                return frame
        raise HTTPException(status_code=422, detail="Word files must contain a table with Date, Product, Quantity, and Revenue (or Unit Price) columns.")
    if extension == ".pdf":
        import pdfplumber

        with pdfplumber.open(BytesIO(raw)) as document:
            if len(document.pages) > 100:
                raise HTTPException(status_code=413, detail="PDF files are limited to 100 pages.")
            for page in document.pages:
                for table in page.extract_tables() or []:
                    frame = _frame_from_table(table)
                    if frame is not None:
                        return frame
        raise HTTPException(status_code=422, detail="PDF must contain a selectable-text sales table. Scanned/image-only PDFs are not supported.")
    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Use CSV, TSV, TXT, JSON, XLS, XLSX, XLSM, DOCX, or a text-based PDF.",
    )


def _normalize_sales_dataframe(df):
    if df.empty:
        raise HTTPException(status_code=422, detail="The uploaded file contains no data rows.")

    mapping = _column_mapping(df.columns)
    missing = [name for name in ("Date", "Product", "Quantity") if name not in mapping]
    if "Revenue" not in mapping and not ("UnitPrice" in mapping and "Quantity" in mapping):
        missing.append("Revenue (or Unit Price)")
    if missing:
        raise HTTPException(
            status_code=422,
            detail=f"Could not find required sales columns: {', '.join(missing)}. Include Date, Product, Quantity, and Revenue; Revenue can be calculated from Quantity × Unit Price.",
        )

    result = df.copy()
    rename = {source: canonical for canonical, source in mapping.items() if canonical != "UnitPrice"}
    result = result.rename(columns=rename)

    def numeric_series(series):
        normalized = series.astype("string").str.strip()
        negatives = normalized.str.match(r"^\(.*\)$", na=False)
        normalized = normalized.str.replace(r"[\s,$£€₹]", "", regex=True)
        normalized.loc[negatives] = "-" + normalized.loc[negatives].str.strip("()")
        return pd.to_numeric(normalized, errors="coerce")

    result["Quantity"] = numeric_series(result["Quantity"])
    if "Revenue" not in result.columns:
        price = numeric_series(result[mapping["UnitPrice"]])
        result["Revenue"] = result["Quantity"] * price
    else:
        result["Revenue"] = numeric_series(result["Revenue"])
    numeric_values = result[["Quantity", "Revenue"]].to_numpy(dtype=np.float64, na_value=np.nan)
    if np.isnan(numeric_values).any() or not np.isfinite(numeric_values).all():
        raise HTTPException(status_code=422, detail="Quantity and Revenue must contain finite numeric values.")

    dates = pd.to_datetime(result["Date"], errors="coerce", dayfirst=True, format="mixed")
    if dates.isna().any():
        raise HTTPException(status_code=422, detail="Date contains values that could not be read as dates.")
    result["Date"] = dates.dt.strftime("%Y-%m-%d")
    products = result["Product"].astype("string").str.strip().fillna("Unspecified")
    products = products.mask(products.eq(""), "Unspecified")
    result["Product"] = products
    return result


def build_business_context(df):
    """Summarize the uploaded sales data once, keeping prompts compact and grounded."""
    revenue = float(df["Revenue"].sum())
    metrics = get_business_metrics(df)
    product_totals = (
        df.groupby("Product", dropna=False)
        .agg(quantity=("Quantity", "sum"), revenue=("Revenue", "sum"))
        .sort_values("revenue", ascending=False)
    )
    products = [
        {
            "name": str(name)[:100],
            "quantity": round(float(row["quantity"]), 2),
            "revenue": round(float(row["revenue"]), 2),
            "revenueSharePct": round(float(row["revenue"]) / revenue * 100, 1) if revenue else 0,
        }
        for name, row in product_totals.head(10).iterrows()
    ]

    dates = pd.to_datetime(df["Date"], errors="coerce")
    valid_dates = dates.dropna()
    monthly_revenue = []
    best_sales_day = None
    if not valid_dates.empty:
        dated = df.loc[valid_dates.index, ["Revenue"]].copy()
        dated["month"] = valid_dates.dt.to_period("M").astype(str)
        monthly = dated.groupby("month")["Revenue"].sum().sort_index()
        monthly_revenue = [
            {"month": month, "revenue": round(float(amount), 2)}
            for month, amount in monthly.tail(24).items()
        ]
        dated["date"] = valid_dates.dt.strftime("%Y-%m-%d")
        daily = dated.groupby("date")["Revenue"].sum()
        if not daily.empty:
            best_date = daily.idxmax()
            best_sales_day = {"date": best_date, "revenue": round(float(daily.loc[best_date]), 2)}

    start_date = valid_dates.min().strftime("%Y-%m-%d") if not valid_dates.empty else None
    end_date = valid_dates.max().strftime("%Y-%m-%d") if not valid_dates.empty else None
    return {
        "rowCount": int(len(df)),
        "period": {"start": start_date, "end": end_date, "distinctDays": int(valid_dates.dt.normalize().nunique())},
        "revenue": round(revenue, 2),
        "estimatedExpenses": round(float(metrics["expenses"]), 2),
        "estimatedProfit": round(float(metrics["profit"]), 2),
        "estimatedCash": round(float(metrics["cash"]), 2),
        "businessHealth": int(metrics["businessHealth"]),
        "riskBand": metrics["risk"],
        "estimationAssumptions": [
            "The uploaded sales file contains sales only; expenses are estimated as 60% of revenue.",
            "Cash is estimated as 80% of estimated profit. Neither estimate is an observed statement balance.",
        ],
        "products": products,
        "monthlyRevenue": monthly_revenue,
        "bestSalesDay": best_sales_day,
        "limitations": "No verified expenses, cash balances, unit costs, customers, inventory levels, or future forecasts are present in the uploaded sales file.",
    }


# ---------------- REQUEST MODELS ----------------

class Question(BaseModel):
    question: str = Field(min_length=1, max_length=2000)


class HiringRequest(BaseModel):
    employees: int = Field(ge=0, le=10000)


class PriceRequest(BaseModel):
    percentage: float = Field(ge=-100, le=1000, allow_inf_nan=False)

class MarketingRequest(BaseModel):
    budget: float = Field(ge=0, le=1_000_000_000, allow_inf_nan=False)

class LoanRequest(BaseModel):
    loanAmount: float = Field(ge=0, le=1_000_000_000, allow_inf_nan=False)
    interestRate: float = Field(ge=0, le=100, allow_inf_nan=False)
    years: int = Field(ge=1, le=100)

class InventoryRequest(BaseModel):
    quantity: int = Field(ge=0, le=10_000_000)
    costPerUnit: float = Field(ge=0, le=1_000_000_000, allow_inf_nan=False)
    expectedProfitPerUnit: float = Field(ge=-1_000_000_000, le=1_000_000_000, allow_inf_nan=False)

class ExpansionRequest(BaseModel):
    expansionCost: float = Field(ge=0, le=1_000_000_000, allow_inf_nan=False)
    expectedRevenue: float = Field(ge=0, le=1_000_000_000, allow_inf_nan=False)
    expectedExpenses: float = Field(ge=0, le=1_000_000_000, allow_inf_nan=False)

# ---------------- HOME ----------------

@app.get("/")
def home():
    return {
        "message": "Welcome to SME Decision Simulator Backend!"
    }


@app.get("/health")
def health():
    return {"status": "ok"}


# ---------------- DASHBOARD ----------------

@app.get("/dashboard")
def dashboard():
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "No sales file uploaded. Please upload a sales file first."
        }

    try:
        metrics = get_business_metrics(uploaded_data)

        if "OrderID" in uploaded_data.columns:
            order_ids = uploaded_data["OrderID"].astype("string").str.strip()
            valid_order_ids = order_ids.notna() & order_ids.ne("")
            total_orders = int(order_ids.loc[valid_order_ids].nunique())
            order_count_label = "Distinct Invoices"
            average_revenue = round(float(uploaded_data.loc[valid_order_ids].assign(_order_id=order_ids[valid_order_ids]).groupby("_order_id")["Revenue"].sum().mean()), 2) if valid_order_ids.any() else None
            average_revenue_sub = "Per invoice"
        else:
            total_orders = len(uploaded_data)
            order_count_label = "Sales Rows"
            average_revenue = round(float(uploaded_data["Revenue"].mean()), 2)
            average_revenue_sub = "Per sales row"
        total_quantity = uploaded_data["Quantity"].sum()

        best_selling_product = (
            uploaded_data.groupby("Product")["Quantity"]
            .sum()
            .idxmax()
        )

        revenue_by_day = uploaded_data.groupby("Date")["Revenue"].sum()
        highest_day = revenue_by_day.idxmax() if not revenue_by_day.empty else None
        lowest_day = revenue_by_day.idxmin() if not revenue_by_day.empty else None

        return {
            "businessHealth": metrics["businessHealth"],
            "revenue": round(float(metrics["revenue"]), 2),
            "expenses": round(float(metrics["expenses"]), 2),
            "profit": round(float(metrics["profit"]), 2),
            "cash": round(float(metrics["cash"]), 2),
            "runway": None,
            "risk": metrics["risk"],
            "revenueBasis": "Net line revenue; includes negative return/cancellation lines" if (uploaded_data["Revenue"] < 0).any() else "Recorded line revenue",

            "totalOrders": total_orders,
            "totalOrdersLabel": order_count_label,
            "totalQuantity": round(float(total_quantity), 2),
            "quantityBasis": "Net units; negative return/cancellation quantities are included" if (uploaded_data["Quantity"] < 0).any() else "Recorded units",
            "bestSellingProduct": best_selling_product,
            "averageRevenue": average_revenue,
            "averageRevenueSub": average_revenue_sub,
            "highestRevenueDay": highest_day,
            "lowestRevenueDay": lowest_day
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail="Failed to compute dashboard metrics.") from e


# ---------------- BUSINESS DATA UPLOAD ----------------

@app.post("/upload")
async def upload_business_file(file: UploadFile = File(...)):
    global uploaded_data, uploaded_business_context

    try:
        extension = os.path.splitext(file.filename or "")[1].casefold()
        if extension not in SUPPORTED_UPLOAD_EXTENSIONS:
            raise HTTPException(status_code=400, detail="Supported files: CSV, TSV, TXT, JSON, XLS, XLSX, XLSM, DOCX, or text-based PDF.")
        raw = await file.read(MAX_UPLOAD_BYTES + 1)
        if len(raw) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail=f"File must be at most {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
        if not raw:
            raise HTTPException(status_code=400, detail="The uploaded file is empty.")
        df = _normalize_sales_dataframe(_read_sales_file(extension, raw))
        if len(df) > MAX_UPLOAD_ROWS:
            raise HTTPException(status_code=422, detail=f"Sales files cannot contain more than {MAX_UPLOAD_ROWS:,} rows.")

        uploaded_data = df
        uploaded_business_context = build_business_context(df)

        return {
            "success": True,
            "message": "Business sales file uploaded successfully!",
            "filename": file.filename,
            "rows": len(uploaded_data),
            "columns": list(uploaded_data.columns),
            "preview": uploaded_data.head(5).to_dict(orient="records")
        }

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Could not read this file. Make sure it is valid, is not password-protected, and contains a sales table with Date, Product, Quantity, and Revenue (or Unit Price) columns.") from exc


# ---------------- CHART DATA ----------------

@app.get("/chart-data")
def chart_data():
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "No sales file uploaded. Please upload a sales file first."
        }

    try:
        revenue_by_date = (
            uploaded_data.groupby("Date")["Revenue"]
            .sum()
            .reset_index()
        )

        quantity_by_product = (
            uploaded_data.groupby("Product")["Quantity"]
            .sum()
            .reset_index()
        )

        revenue_by_product = (
            uploaded_data.groupby("Product")["Revenue"]
            .sum()
            .reset_index()
        )

        return {
            "revenueByDate": revenue_by_date.to_dict(orient="records"),
            "quantityByProduct": quantity_by_product.to_dict(orient="records"),
            "revenueByProduct": revenue_by_product.to_dict(orient="records")
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail="Failed to compute chart data.") from e


# ---------------- AI BUSINESS ADVISOR ----------------

@app.post("/ask")
def ask(question: Question):
    global uploaded_data, uploaded_business_context

    if uploaded_data is None:
        return {
            "error": "Please upload a CSV before asking AI questions."
        }

    if uploaded_business_context is None:
        uploaded_business_context = build_business_context(uploaded_data)
    answer = ask_business_question(question.question, uploaded_business_context)
    return {"answer": answer}
# ---------------- HIRING SIMULATION ----------------

@app.post("/simulate/hiring")
def hiring_simulation(data: HiringRequest):
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "Please upload a CSV first."
        }

    metrics = get_business_metrics(uploaded_data)

    revenue = metrics["revenue"]
    expenses = metrics["expenses"]
    profit = metrics["profit"]
    cash = metrics["cash"]

    simulation = simulate_hiring(
        revenue=revenue,
        expenses=expenses,
        cash=cash,
        employees=data.employees
    )

    prompt = f"""You are an experienced SME Business Consultant.

Current Business

Revenue: ₹{int(revenue)}
Expenses: ₹{int(expenses)}
Profit: ₹{int(profit)}
Cash: ₹{int(cash)}

Hiring Simulation

Employees to Hire: {data.employees}
Annual Salary Cost: ₹{simulation['annualSalaryCost']}
Projected Expenses: ₹{simulation['projectedExpenses']}
Projected Profit: ₹{simulation['projectedProfit']}
Projected Cash: ₹{simulation['projectedCash']}
Business Health: {simulation['businessHealth']}%
Risk Level: {simulation['risk']}

Should the company hire these employees?

Respond ONLY in the following format.

Decision:
✅ Approve
or
❌ Reject

Reason:
Maximum 2 short sentences.

Recommendations:
- Recommendation 1
- Recommendation 2
- Recommendation 3

Rules:
- Maximum 80 words.
- Focus only on the hiring decision.
- No introduction.
- No conclusion.
- Do not explain calculations.
"""

    ai_answer = ask_ai(prompt)

    return {
        "currentBusiness": {
            "revenue": int(revenue),
            "expenses": int(expenses),
            "profit": int(profit),
            "cash": int(cash)
        },
        "simulation": simulation,
        "aiRecommendation": ai_answer
    }

# ---------------- PRICE SIMULATION ----------------

@app.post("/simulate/price")
def price_simulation(data: PriceRequest):
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "Please upload a CSV first."
        }

    metrics = get_business_metrics(uploaded_data)

    simulation = simulate_price_change(
        revenue=metrics["revenue"],
        expenses=metrics["expenses"],
        percentage=data.percentage
    )

    prompt = f"""You are an experienced SME Business Consultant.

Current Business

Revenue: ₹{int(metrics['revenue'])}
Expenses: ₹{int(metrics['expenses'])}
Profit: ₹{int(metrics['profit'])}
Cash: ₹{int(metrics['cash'])}

Price Change Simulation

Price Change: {data.percentage}%
Projected Revenue: ₹{simulation['newRevenue']}
Projected Profit: ₹{simulation['newProfit']}
Projected Cash: ₹{simulation['newCash']}
Business Health: {simulation['businessHealth']}%
Risk Level: {simulation['risk']}

Should the company increase prices?

Respond ONLY in the following format.

Decision:
✅ Approve
or
❌ Reject

Reason:
Maximum 2 short sentences.

Recommendations:
- Recommendation 1
- Recommendation 2
- Recommendation 3

Rules:
- Maximum 80 words.
- Focus only on the pricing decision.
- No introduction.
- No conclusion.
- Do not explain calculations.
"""

    ai_answer = ask_ai(prompt)

    return {
        "currentBusiness": {
            "revenue": int(metrics["revenue"]),
            "expenses": int(metrics["expenses"]),
            "profit": int(metrics["profit"]),
            "cash": int(metrics["cash"])
        },
        "simulation": simulation,
        "aiRecommendation": ai_answer
    }

# ---------------- MARKETING SIMULATION ----------------

@app.post("/simulate/marketing")
def marketing_simulation(data: MarketingRequest):
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "Please upload a CSV first."
        }

    metrics = get_business_metrics(uploaded_data)

    simulation = simulate_marketing(
        revenue=metrics["revenue"],
        expenses=metrics["expenses"],
        budget=data.budget
    )

    prompt = f"""
You are an experienced SME Business Consultant.

Current Business

Revenue: ₹{int(metrics['revenue'])}
Expenses: ₹{int(metrics['expenses'])}
Profit: ₹{int(metrics['profit'])}
Cash: ₹{int(metrics['cash'])}

Marketing Simulation

Marketing Budget: ₹{simulation['marketingBudget']}
Estimated Revenue Increase: ₹{simulation['estimatedRevenueIncrease']}
Projected Revenue: ₹{simulation['newRevenue']}
Projected Profit: ₹{simulation['newProfit']}
Projected Cash: ₹{simulation['newCash']}
Business Health: {simulation['businessHealth']}%
Risk Level: {simulation['risk']}

Should the company invest this amount in marketing?

Respond ONLY in the following format.

Decision:
Approve or Reject

Reason:
Maximum 2 short sentences.

Recommendations:
• Recommendation 1
• Recommendation 2
• Recommendation 3

Rules:
- Maximum 80 words.
- Focus only on the marketing investment decision.
- No introduction.
- No conclusion.
- Do not explain calculations.
"""

    ai_answer = ask_ai(prompt)

    return {
        "currentBusiness": {
            "revenue": int(metrics["revenue"]),
            "expenses": int(metrics["expenses"]),
            "profit": int(metrics["profit"]),
            "cash": int(metrics["cash"])
        },
        "simulation": simulation,
        "aiRecommendation": ai_answer
    }

# ---------------- LOAN SIMULATION ----------------

@app.post("/simulate/loan")
def loan_simulation(data: LoanRequest):
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "Please upload a CSV first."
        }

    metrics = get_business_metrics(uploaded_data)

    simulation = simulate_loan(
        cash=metrics["cash"],
        loan_amount=data.loanAmount,
        interest_rate=data.interestRate,
        years=data.years,
        revenue=metrics["revenue"],
        profit=metrics["profit"]
    )

    prompt = f"""
You are an experienced SME Business Consultant.

Current Business

Cash Available: ₹{int(metrics['cash'])}
Business Health: {metrics['businessHealth']}%
Risk Level: {metrics['risk']}

Loan Simulation

Loan Amount: ₹{simulation['loanAmount']}
Interest Rate: {simulation['interestRate']}%
Loan Period: {simulation['loanPeriod']} years
Total Interest: ₹{simulation['totalInterest']}
Total Repayment: ₹{simulation['totalRepayment']}
Monthly EMI: ₹{simulation['monthlyEMI']}
Cash After Loan: ₹{simulation['newCash']}
Cash After First EMI: ₹{simulation['cashAfterFirstEMI']}
Business Health: {simulation['businessHealth']}%
Risk Level: {simulation['risk']}

Should the company take this loan?

Respond ONLY in the following format.

Decision:
Approve or Reject

Reason:
Maximum 2 short sentences.

Recommendations:
• Recommendation 1
• Recommendation 2
• Recommendation 3

Rules:
- Maximum 80 words.
- Focus only on the loan decision.
- No introduction.
- No conclusion.
- Do not explain calculations.
"""

    ai_answer = ask_ai(prompt)

    return {
        "currentBusiness": {
            "cash": int(metrics["cash"]),
            "businessHealth": metrics["businessHealth"],
            "risk": metrics["risk"]
        },
        "simulation": simulation,
        "aiRecommendation": ai_answer
    }

# ---------------- INVENTORY SIMULATION ----------------

@app.post("/simulate/inventory")
def inventory_simulation(data: InventoryRequest):
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "Please upload a CSV first."
        }

    metrics = get_business_metrics(uploaded_data)

    simulation = simulate_inventory(
        cash=metrics["cash"],
        quantity=data.quantity,
        cost_per_unit=data.costPerUnit,
        expected_profit_per_unit=data.expectedProfitPerUnit,
        revenue=metrics["revenue"]
    )

    prompt = f"""
You are an experienced SME Business Consultant.

Current Business

Cash Available: ₹{int(metrics['cash'])}
Business Health: {metrics['businessHealth']}%
Risk Level: {metrics['risk']}

Inventory Simulation

Quantity: {simulation['quantity']}
Cost Per Unit: ₹{simulation['costPerUnit']}
Inventory Cost: ₹{simulation['inventoryCost']}
Expected Profit: ₹{simulation['expectedProfit']}
Remaining Cash: ₹{simulation['remainingCash']}
Projected Cash: ₹{simulation['projectedCash']}
Business Health: {simulation['businessHealth']}%
Risk Level: {simulation['risk']}

Should the company purchase this inventory?

Respond ONLY in the following format.

Decision:
Approve or Reject

Reason:
Maximum 2 short sentences.

Recommendations:
• Recommendation 1
• Recommendation 2
• Recommendation 3

Rules:
- Maximum 80 words.
- Focus only on the inventory purchase decision.
- No introduction.
- No conclusion.
- Do not explain calculations.
"""

    ai_answer = ask_ai(prompt)

    return {
        "currentBusiness": {
            "cash": int(metrics["cash"]),
            "businessHealth": metrics["businessHealth"],
            "risk": metrics["risk"]
        },
        "simulation": simulation,
        "aiRecommendation": ai_answer
    }

# ---------------- EXPANSION SIMULATION ----------------

@app.post("/simulate/expansion")
def expansion_simulation(data: ExpansionRequest):
    global uploaded_data

    if uploaded_data is None:
        return {
            "error": "Please upload a CSV first."
        }

    metrics = get_business_metrics(uploaded_data)

    simulation = simulate_expansion(
        cash=metrics["cash"],
        expansion_cost=data.expansionCost,
        expected_revenue=data.expectedRevenue,
        expected_expenses=data.expectedExpenses
    )

    prompt = f"""
You are an experienced SME Business Consultant.

Current Business

Cash Available: ₹{int(metrics['cash'])}
Business Health: {metrics['businessHealth']}%
Risk Level: {metrics['risk']}

Expansion Simulation

Expansion Cost: ₹{simulation['expansionCost']}
Expected Revenue: ₹{simulation['expectedRevenue']}
Expected Expenses: ₹{simulation['expectedExpenses']}
Projected Profit: ₹{simulation['projectedProfit']}
Remaining Cash: ₹{simulation['remainingCash']}
Projected Cash: ₹{simulation['projectedCash']}
Business Health: {simulation['businessHealth']}%
Risk Level: {simulation['risk']}

Should the company proceed with this expansion?

Respond ONLY in the following format.

Decision:
Approve or Reject

Reason:
Maximum 2 short sentences.

Recommendations:
• Recommendation 1
• Recommendation 2
• Recommendation 3

Rules:
- Maximum 80 words.
- Focus only on the expansion decision.
- No introduction.
- No conclusion.
- Do not explain calculations.
"""

    ai_answer = ask_ai(prompt)

    return {
        "currentBusiness": {
            "cash": int(metrics["cash"]),
            "businessHealth": metrics["businessHealth"],
            "risk": metrics["risk"]
        },
        "simulation": simulation,
        "aiRecommendation": ai_answer
    }
# ---------------- REPORT DOWNLOADS ----------------

@app.get("/reports/{report_type}/pdf")
def download_report_pdf(report_type: str):
    global uploaded_data

    if uploaded_data is None:
        return {"error": "Please upload a CSV first."}

    if report_type not in REPORT_TYPES:
        return {"error": f"Unknown report type: {report_type}"}

    try:
        data = build_report_data(uploaded_data)
        pdf_buffer = generate_pdf_report(report_type, data)
        filename = f"{REPORT_TYPES[report_type]['title'].replace(' ', '_')}.pdf"

        return StreamingResponse(
            pdf_buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail="Failed to generate PDF report.") from e


@app.get("/reports/{report_type}/csv")
def download_report_csv(report_type: str):
    global uploaded_data

    if uploaded_data is None:
        return {"error": "Please upload a CSV first."}

    if report_type not in REPORT_TYPES:
        return {"error": f"Unknown report type: {report_type}"}

    try:
        data = build_report_data(uploaded_data)
        csv_buffer = generate_csv_report(report_type, data)
        filename = f"{REPORT_TYPES[report_type]['title'].replace(' ', '_')}.csv"

        return StreamingResponse(
            csv_buffer,
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail="Failed to generate CSV report.") from e


@app.get("/reports/export-all")
def export_all_reports():
    global uploaded_data

    if uploaded_data is None:
        return {"error": "Please upload a CSV first."}

    try:
        zip_buffer = generate_all_reports_zip(uploaded_data)

        return StreamingResponse(
            zip_buffer,
            media_type="application/zip",
            headers={"Content-Disposition": 'attachment; filename="SME_Reports.zip"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail="Failed to export reports.") from e
